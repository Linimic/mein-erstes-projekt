# Erzeugt Dienstdoku_Mitarbeiter_2026.xlsx
#   python3 erstelle_dienstdoku_mitarbeiter.py                       -> leere Vorlage
#   python3 erstelle_dienstdoku_mitarbeiter.py daten.json ziel.xlsx  -> Vorlage mit übernommenen Einträgen
import sys, json, datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side, Protection
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import DataPoint
from openpyxl.workbook.properties import CalcProperties

OUT = "Dienstdoku_Mitarbeiter_2026.xlsx"
JAHR, EINRICHTUNG = 2026, "Villa K."
ANGEBOTE = ["MITTAG", "JUGENDCAFE", "JUGENDCAFE NEXT", "OSA", "VERMIETUNGEN", "TRÄNKEWEG", "TANZEN",
            "KONZERTE", "JUNGS", "MÄDCHEN", "JSA", "RADIKALISIERUNG", "SCHULE", "WORKSHOPS", "CARITAS",
            "KINDERTHEATER", "WOODROCK", "KOOPERATIONEN", "ÖA", "OUTDOORFUN", "ONLINE", "MOJA",
            "SKATERPLATZ", "BERATUNG / TEL"]
MONATSNAMEN = ["Jänner", "Februar", "März", "April", "Mai", "Juni", "Juli", "August", "September",
               "Oktober", "November", "Dezember"]
NA, NT = 30, 30            # Plätze für Angebote / Team
FIRST, LAST = 6, 3005      # Datenzeilen ERFASSUNG
KFIRST, KLAST = 5, 370     # Kalenderzeilen (366 Tage)

# ---------- Stil ----------
F = "Arial"
NAVY, BLUE, CALC, YEL, ZEBRA = "1F3864", "DDEBF7", "F2F2F2", "FFF2CC", "F7F9FC"
GCOL = ["8E6CC4", "F2A541", "3E8EC2"]                          # w, d, m
ACOL = ["C9DDF0", "9CC3E6", "5B9BD5", "2E75B6", "1F4E79"]      # unter 10 … ab 18
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
UNLOCK = Protection(locked=False)


def font(**k):
    k.setdefault("name", F)
    return Font(**k)


def fill(c):
    return PatternFill("solid", fgColor=c)


def hdr(c, color=NAVY):
    c.font = font(bold=True, color="FFFFFF", size=9)
    c.fill = fill(color)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    c.border = BOX


def cell(ws, ref, value=None, bold=False, size=9, bg=CALC, fmt=None, center=True, color="000000"):
    c = ws[ref]
    if value is not None:
        c.value = value
    c.font = font(bold=bold, size=size, color=color)
    if bg:
        c.fill = fill(bg)
    c.border = BOX
    if center:
        c.alignment = Alignment(horizontal="center", vertical="center")
    if fmt:
        c.number_format = fmt
    return c


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = font(size=16, bold=True, color=NAVY)
    ws["A2"] = sub
    ws["A2"].font = font(italic=True, size=9, color="595959")


def section(ws, ref, text):
    ws[ref] = text
    ws[ref].font = font(size=11, bold=True, color=NAVY)


def tile(ws, row, col, label, formula, fmt="0"):
    ws.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col + 1)
    ws.merge_cells(start_row=row + 1, start_column=col, end_row=row + 1, end_column=col + 1)
    a = ws.cell(row, col, label)
    a.font = font(size=8, color="595959")
    a.alignment = Alignment(horizontal="center", vertical="bottom", wrap_text=True)
    v = ws.cell(row + 1, col, formula)
    v.font = font(size=16, bold=True, color=NAVY)
    v.alignment = Alignment(horizontal="center", vertical="center")
    v.number_format = fmt
    for r in (row, row + 1):
        for c in (col, col + 1):
            ws.cell(r, c).fill = fill(BLUE)
            ws.cell(r, c).border = BOX
    ws.row_dimensions[row + 1].height = 28


def protect(ws):
    p = ws.protection
    p.sheet = True
    p.autoFilter = False
    p.sort = False
    p.formatColumns = False
    p.formatRows = False
    p.selectLockedCells = False
    p.selectUnlockedCells = False
    p.objects = False      # Diagramme bleiben kopierbar


def labels(pct=False):
    d = DataLabelList()
    d.showVal = not pct
    d.showPercent = pct
    d.showSerName = d.showCatName = d.showLegendKey = d.showLeaderLines = False
    return d


def color_series(ch, cols):
    for s, c in zip(ch.series, cols):
        s.graphicalProperties.solidFill = c
        s.graphicalProperties.line.solidFill = c


def color_points(ch, cols):
    for i, c in enumerate(cols):
        pt = DataPoint(idx=i)
        pt.graphicalProperties.solidFill = c
        ch.series[0].dPt.append(pt)


# ---------- Bezüge ----------
Y = "LISTEN!$F$2"
ORG = "LISTEN!$F$3"
ANG = f"LISTEN!$A$2:$A${NA + 1}"
TEAM = f"LISTEN!$C$2:$C${NT + 1}"
MON = "LISTEN!$H$2:$H$13"


def R(col):   # Spalte der Erfassung
    return f"ERFASSUNG!${col}${FIRST}:${col}${LAST}"


def K(col):   # Spalte des Kalenders
    return f"KALENDER!${col}${KFIRST}:${col}${KLAST}"


def iso(x):   # ISO-Kalenderwoche ohne neuere Excel-Funktionen
    return (f"INT(({x}-DATE(YEAR({x}-WEEKDAY({x}-1)+4),1,3)"
            f"+WEEKDAY(DATE(YEAR({x}-WEEKDAY({x}-1)+4),1,3))+5)/7)")


WTAG = '"Mo","Di","Mi","Do","Fr","Sa","So"'
# Kennzahlen: Name, Spalte Erfassung, Spalte Kalender
MET = [("Stunden", "E", "H"), ("w", "F", "I"), ("d", "G", "J"), ("m", "H", "K"), ("gesamt", "I", "L"),
       ("unter 10", "J", "M"), ("ab 10", "K", "N"), ("ab 14", "L", "O"), ("ab 16", "M", "P"), ("ab 18", "N", "Q")]
AGES = ["unter 10", "ab 10", "ab 14", "ab 16", "ab 18"]

wb = Workbook()

# =====================================================================
# ANLEITUNG
# =====================================================================
a = wb.active
a.title = "ANLEITUNG"
a.sheet_view.showGridLines = False
a.column_dimensions["A"].width = 3
a.column_dimensions["B"].width = 16
for i in range(3, 23):
    a.column_dimensions[L(i)].width = 11
a.column_dimensions["C"].width = 16
a.column_dimensions["T"].width = 18
a.column_dimensions["U"].width = 22
a["B2"] = "Dienstdokumentation – so funktioniert's"
a["B2"].font = font(size=18, bold=True, color=NAVY)
steps = [
    ("EINMALIG", "LISTEN", "Jahr, Einrichtung, Angebote und Namen des Teams eintragen (gelbe Zellen). "
                           "Neue Angebote/Namen erscheinen sofort in allen Dropdowns und Auswertungen."),
    ("TÄGLICH", "ERFASSUNG", "Pro durchgeführtem Angebot EINE Zeile: Datum, Angebot, Stunden, Anzahl w/d/m, "
                             "Altersgruppen, wer Dienst hatte und Notizen. Einfach unten weiterschreiben."),
    ("WÖCHENTLICH", "WOCHE", "Kalenderwoche oben eintragen → Tage, Angebote, Team-Einsätze und "
                             "'Wichtiges der Woche' auf einen Blick (z. B. für die Teamsitzung)."),
    ("MONATLICH", "MONAT", "Monat oben auswählen → Monatszusammenfassung mit Kennzahlen, Tabellen und Diagrammen."),
    ("JÄHRLICH", "JAHRESBERICHT", "Alles für den Jahresbericht: Kennzahlen mit Vorjahresvergleich, fertiger "
                                  "Textbaustein, Monats-, Angebots-, Wochentags- und Wochenstatistik, Diagramme."),
]
r = 4
for when, sheet, txt in steps:
    c = a.cell(r, 2, when)
    c.font = font(bold=True, color="FFFFFF", size=9)
    c.fill = fill(NAVY)
    c.alignment = Alignment(horizontal="center", vertical="center")
    c = a.cell(r, 3, sheet)
    c.font = font(bold=True, color=NAVY, size=10)
    c.fill = fill(BLUE)
    c.alignment = Alignment(horizontal="center", vertical="center")
    a.merge_cells(start_row=r, start_column=4, end_row=r, end_column=16)
    c = a.cell(r, 4, txt)
    c.font = font(size=10)
    c.alignment = Alignment(wrap_text=True, vertical="center")
    a.row_dimensions[r].height = 30
    r += 1

r += 1
section(a, f"B{r}", "Beispiel – so sieht eine Zeile in ERFASSUNG aus")
r += 1
EX_HEAD = ["Datum", "Tag", "KW", "Angebot", "Stunden", "w", "d", "m", "gesamt", "unter 10", "ab 10", "ab 14",
           "ab 16", "ab 18", "Alters-Check", "Dienst 1", "Dienst 2", "Dienst 3", "Was war los?",
           "Wichtig / Vorkommnisse"]
EX_ROW = ["22.09.2026", "Di", 39, "JUGENDCAFE", 4, 7, 1, 12, 20, 0, 5, 9, 4, 2, "✓", "Anna", "Max", "",
          "Billard-Turnier", "Fenster WC defekt"]
for i, (h, v) in enumerate(zip(EX_HEAD, EX_ROW)):
    hdr(a.cell(r, 2 + i, h))
    c = a.cell(r + 1, 2 + i, v)
    c.font = font(size=9)
    c.border = BOX
    c.alignment = Alignment(horizontal="center")
    if h in ("Tag", "KW", "gesamt", "Alters-Check"):
        c.fill = fill(CALC)
a.row_dimensions[r].height = 30
r += 3

section(a, f"B{r}", "Farben")
r += 1
for bg, txt in [(YEL, "Auswahl / Einstellung (z. B. Kalenderwoche, Monat, Jahr, Listen)"),
                ("FFFFFF", "Eingabe"),
                (CALC, "rechnet automatisch – nicht überschreiben (Blätter sind ohne Passwort geschützt: "
                       "Überprüfen → Blattschutz aufheben)"),
                ("F8CBAD", "Alters-Check passt nicht: Summe der Altersgruppen ≠ gesamt")]:
    c = a.cell(r, 2)
    c.fill = fill(bg)
    c.border = BOX
    a.cell(r, 3, txt).font = font(size=10)
    r += 1
r += 1
section(a, f"B{r}", "Begriffe")
r += 1
for k, v in [("Kontakte", "Anzahl Besuche. Wer an 3 Tagen kommt, zählt 3 Kontakte."),
             ("Einträge", "Wie oft ein Angebot stattgefunden hat (= Anzahl Zeilen)."),
             ("Öffnungstage", "Tage mit mindestens einem Eintrag."),
             ("Altersgruppen", "unter 10 · ab 10 = 10–13 · ab 14 = 14–15 · ab 16 = 16–17 · ab 18 = 18+"),
             ("Team-Stunden", "Stunden des Angebots, bei dem die Person im Dienst eingetragen ist.")]:
    a.cell(r, 2, k).font = font(bold=True, size=10)
    a.cell(r, 3, v).font = font(size=10)
    r += 1

# =====================================================================
# LISTEN
# =====================================================================
li = wb.create_sheet("LISTEN")
li.sheet_view.showGridLines = False
for col, w in zip("ABCDEFGH", [26, 3, 22, 3, 14, 16, 3, 14]):
    li.column_dimensions[col].width = w
for ref, t in [("A1", "Angebote"), ("C1", "Team (Namen)"), ("E1", "Einstellung"), ("F1", "Wert"),
               ("H1", "Monate")]:
    hdr(li[ref])
for i in range(NA):
    cell(li, f"A{2 + i}", ANGEBOTE[i] if i < len(ANGEBOTE) else None, bg=YEL, center=False, size=10)
for i in range(NT):
    cell(li, f"C{2 + i}", None, bg=YEL, center=False, size=10)
cell(li, "E2", "Jahr", bg=None, center=False, size=10)
cell(li, "F2", JAHR, bold=True, bg=YEL, size=10)
cell(li, "E3", "Einrichtung", bg=None, center=False, size=10)
cell(li, "F3", EINRICHTUNG, bold=True, bg=YEL, size=10)
for i, m in enumerate(MONATSNAMEN):
    cell(li, f"H{2 + i}", m, bg=CALC, center=False, size=10)
li["E6"] = "Gelbe Zellen anpassen. Neue Einträge einfach in freie Zeilen schreiben."
li["E6"].font = font(italic=True, size=9, color="595959")

# =====================================================================
# ERFASSUNG
# =====================================================================
er = wb.create_sheet("ERFASSUNG")
er.sheet_view.showGridLines = False
title(er, "Tägliche Erfassung",
      "Pro durchgeführtem Angebot eine Zeile. Weiß = eingeben, grau = rechnet automatisch. "
      "Dropdowns für Angebot und Dienst kommen aus dem Blatt LISTEN.")
er["A1"] = f'="Tägliche Erfassung "&{Y}&" – "&{ORG}'
WID = [11, 5, 5, 20, 8, 6, 6, 6, 8, 7, 7, 7, 7, 7, 11, 13, 13, 13, 34, 34]
for i, w in enumerate(WID):
    er.column_dimensions[L(i + 1)].width = w
er["D4"] = "SUMME (Filter beachtet)"
er["D4"].font = font(bold=True, size=9)
er["D4"].alignment = Alignment(horizontal="right")
for col in "EFGHIJKLMN":
    cell(er, f"{col}4", f"=SUBTOTAL(9,{col}{FIRST}:{col}{LAST})", bold=True, bg=BLUE)
for i, h in enumerate(EX_HEAD):
    hdr(er.cell(5, i + 1, h))
er.row_dimensions[5].height = 30
# Hilfsspalten (ausgeblendet)
for col, h in [("V", "Monat"), ("W", "lfd Wichtig Woche"), ("X", "lfd Wichtig Monat")]:
    er[f"{col}5"] = h
    er.column_dimensions[col].hidden = True

INPUT_COLS = set("ADEFGHJKLMNPQRST")
for r in range(FIRST, LAST + 1):
    er[f"B{r}"] = f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),{WTAG}))'
    er[f"C{r}"] = f'=IF(A{r}="","",{iso(f"A{r}")})'
    er[f"I{r}"] = f'=IF(COUNT(F{r}:H{r})=0,"",SUM(F{r}:H{r}))'
    er[f"O{r}"] = (f'=IF(OR(I{r}="",COUNT(J{r}:N{r})=0),"",IF(SUM(J{r}:N{r})=I{r},"✓",'
                   f'"⚠ "&SUM(J{r}:N{r})&" statt "&I{r}))')
    er[f"V{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    er[f"W{r}"] = f'=N(W{r - 1})+IF(AND(C{r}<>"",T{r}<>""),IF(C{r}=WOCHE!$B$3,1,0),0)'
    er[f"X{r}"] = f'=N(X{r - 1})+IF(AND(V{r}<>"",T{r}<>""),IF(V{r}=MONAT!$AZ$3,1,0),0)'
    for ci in range(1, 21):
        col = L(ci)
        c = er.cell(r, ci)
        c.border = BOX
        c.font = font(size=10)
        if col in INPUT_COLS:
            c.protection = UNLOCK
        else:
            c.fill = fill(CALC)
        if col in "ST":
            c.alignment = Alignment(wrap_text=True, vertical="top")
        elif col != "D" and col not in "PQR":
            c.alignment = Alignment(horizontal="center")
    er[f"A{r}"].number_format = "DD.MM.YYYY"
    er[f"E{r}"].number_format = "0.0#"
er.freeze_panes = f"E{FIRST}"
er.auto_filter.ref = f"A5:T{LAST}"

dv = DataValidation(type="date", operator="between", formula1=f"DATE({Y},1,1)", formula2=f"DATE({Y},12,31)",
                    showErrorMessage=True, errorTitle="Datum", error="Bitte ein Datum im eingestellten Jahr (TT.MM.JJJJ).")
dv.add(f"A{FIRST}:A{LAST}")
dv2 = DataValidation(type="list", formula1=ANG, showErrorMessage=True, errorTitle="Angebot",
                     error="Bitte Angebot aus der Liste wählen – neue Angebote im Blatt LISTEN ergänzen.")
dv2.add(f"D{FIRST}:D{LAST}")
dv3 = DataValidation(type="list", formula1=TEAM, showErrorMessage=True, errorTitle="Dienst",
                     error="Bitte Namen aus der Liste wählen – neue Namen im Blatt LISTEN ergänzen.")
dv3.add(f"P{FIRST}:R{LAST}")
dv4 = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True,
                     error="Bitte eine ganze Zahl ≥ 0 eingeben.")
dv4.add(f"F{FIRST}:H{LAST}")
dv4.add(f"J{FIRST}:N{LAST}")
dv5 = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", showErrorMessage=True,
                     error="Stunden zwischen 0 und 24 (z. B. 3,5).")
dv5.add(f"E{FIRST}:E{LAST}")
for d in (dv, dv2, dv3, dv4, dv5):
    er.add_data_validation(d)
# Tage abwechselnd schattiert, Alters-Check rot
er.conditional_formatting.add(f"A{FIRST}:N{LAST}", FormulaRule(
    formula=[f"AND(ISNUMBER($A{FIRST}),MOD($A{FIRST},2)=1)"], fill=fill("EAF1FB")))
er.conditional_formatting.add(f"O{FIRST}:O{LAST}", FormulaRule(
    formula=[f'LEFT($O{FIRST},1)="⚠"'], fill=fill("F8CBAD"), font=Font(color="9C0006", bold=True)))
protect(er)

# =====================================================================
# KALENDER (Hilfsblatt: jeder Tag des Jahres)
# =====================================================================
ka = wb.create_sheet("KALENDER")
ka.sheet_view.showGridLines = False
title(ka, "Kalender (rechnet automatisch)", "Hilfsblatt: jeder Tag des Jahres mit Summen aus der Erfassung.")
KH = ["Datum", "Tag", "KW", "Monat", "Wochentag-Nr", "Einträge", "offen", "Stunden", "w", "d", "m", "gesamt",
      "unter 10", "ab 10", "ab 14", "ab 16", "ab 18"]
for i, h in enumerate(KH):
    hdr(ka.cell(4, i + 1, h))
    ka.column_dimensions[L(i + 1)].width = 11 if i == 0 else 8
for r in range(KFIRST, KLAST + 1):
    i = r - KFIRST
    ka[f"A{r}"] = f'=IF(YEAR(DATE({Y},1,1)+{i})={Y},DATE({Y},1,1)+{i},"")'
    ka[f"B{r}"] = f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),{WTAG}))'
    ka[f"C{r}"] = f'=IF(A{r}="","",{iso(f"A{r}")})'
    ka[f"D{r}"] = f'=IF(A{r}="","",MONTH(A{r}))'
    ka[f"E{r}"] = f'=IF(A{r}="","",WEEKDAY(A{r},2))'
    ka[f"F{r}"] = f'=IF(A{r}="","",COUNTIF({R("A")},A{r}))'
    ka[f"G{r}"] = f'=IF(A{r}="","",IF(F{r}>0,1,0))'
    for name, ec, kc in MET:
        ka[f"{kc}{r}"] = f'=IF(A{r}="","",SUMIFS({R(ec)},{R("A")},A{r}))'
    ka[f"A{r}"].number_format = "DD.MM.YYYY"
    for ci in range(1, 18):
        ka.cell(r, ci).font = font(size=9)
        ka.cell(r, ci).alignment = Alignment(horizontal="center")
ka.freeze_panes = "B5"
protect(ka)


# =====================================================================
# gemeinsame Bausteine für WOCHE / MONAT / JAHR
# =====================================================================
def tiles(ws, row, f_open, f_entries, f_hours, f_w, f_d, f_m):
    kont = f"({f_w}+{f_d}+{f_m})"
    tile(ws, row, 1, "Öffnungstage", f"={f_open}")
    tile(ws, row, 3, "Einträge (Angebote)", f"={f_entries}")
    tile(ws, row, 5, "Kontakte gesamt", f"={kont}")
    tile(ws, row, 7, "Ø Kontakte pro Öffnungstag", f"=IF(N({f_open})=0,0,{kont}/{f_open})", "0.0")
    tile(ws, row, 9, "Stunden", f"={f_hours}", "0.#")
    tile(ws, row, 11, "weiblich", f"=IF({kont}=0,0,{f_w}/{kont})", "0%")
    tile(ws, row, 13, "divers", f"=IF({kont}=0,0,{f_d}/{kont})", "0%")
    tile(ws, row, 15, "männlich", f"=IF({kont}=0,0,{f_m}/{kont})", "0%")


def angebot_table(ws, top, crit):
    """Tabelle je Angebot. crit = zusätzliche SUMIFS-Kriterien (',Bereich,Wert') oder ''."""
    heads = ["Angebot", "Einträge"] + [m[0] for m in MET] + ["Anteil"]
    for i, h in enumerate(heads):
        hdr(ws.cell(top, 1 + i, h))
    first, last = top + 1, top + NA
    for k in range(NA):
        r = first + k
        nm = f"$A{r}"
        ws[f"A{r}"] = f'=IF(LISTEN!$A${2 + k}="","",LISTEN!$A${2 + k})'
        ws[f"B{r}"] = f'=IF({nm}="","",COUNTIFS({R("D")},{nm}{crit}))'
        for j, (name, ec, kc) in enumerate(MET):
            ws.cell(r, 3 + j, f'=IF({nm}="","",SUMIFS({R(ec)},{R("D")},{nm}{crit}))')
        ws.cell(r, 13, f'=IF(OR({nm}="",N($G${last + 1})=0),"",G{r}/$G${last + 1})').number_format = "0%"
        for c in range(1, 14):
            x = ws.cell(r, c)
            x.border = BOX
            x.font = font(size=9)
            if c > 1:
                x.alignment = Alignment(horizontal="center")
            if k % 2:
                x.fill = fill(ZEBRA)
    ws.cell(last + 1, 1, "SUMME")
    for c in range(2, 13):
        ws.cell(last + 1, c, f"=SUM({L(c)}{first}:{L(c)}{last})")
    for c in range(1, 14):
        cell(ws, f"{L(c)}{last + 1}", bold=True, bg=BLUE, center=c > 1)
    return first, last


def team_table(ws, top, col, crit):
    c0 = col
    for i, h in enumerate(["Team", "Einsätze", "Stunden"]):
        hdr(ws.cell(top, c0 + i, h))
    for k in range(NT):
        r = top + 1 + k
        nm = f"${L(c0)}{r}"
        ws.cell(r, c0, f'=IF(LISTEN!$C${2 + k}="","",LISTEN!$C${2 + k})')
        cnt = "+".join(f'COUNTIFS({R(x)},{nm}{crit})' for x in "PQR")
        hrs = "+".join(f'SUMIFS({R("E")},{R(x)},{nm}{crit})' for x in "PQR")
        ws.cell(r, c0 + 1, f'=IF({nm}="","",{cnt})')
        ws.cell(r, c0 + 2, f'=IF({nm}="","",{hrs})').number_format = "0.#"
        for c in range(c0, c0 + 3):
            x = ws.cell(r, c)
            x.border = BOX
            x.font = font(size=9)
            if c > c0:
                x.alignment = Alignment(horizontal="center")
            if k % 2:
                x.fill = fill(ZEBRA)


def wichtig_list(ws, top, col, n, cumcol, helpcol):
    heads = ["Datum", "Angebot", "Dienst", "Wichtig / Vorkommnisse"]
    spans = [1, 2, 2, 6]
    c = col
    for h, s in zip(heads, spans):
        if s > 1:
            ws.merge_cells(start_row=top, start_column=c, end_row=top, end_column=c + s - 1)
        for cc in range(c, c + s):
            hdr(ws.cell(top, cc, h if cc == c else None))
        c += s
    for k in range(n):
        r = top + 1 + k
        ws[f"{helpcol}{r}"] = f'=IFERROR(MATCH({k + 1},{R(cumcol)},0),"")'
        c = col
        for (src, s) in zip("ADPT", spans):
            if s > 1:
                ws.merge_cells(start_row=r, start_column=c, end_row=r, end_column=c + s - 1)
            ws.cell(r, c, f'=IF(${helpcol}{r}="","",INDEX({R(src)},${helpcol}{r}))')
            for cc in range(c, c + s):
                x = ws.cell(r, cc)
                x.border = BOX
                x.font = font(size=9)
                x.alignment = Alignment(vertical="top", wrap_text=(src == "T"))
            c += s
        ws.cell(r, col).number_format = "DD.MM."
    ws.column_dimensions[helpcol].hidden = True


def gender_age_block(ws, top, fw, fd, fm, fages):
    heads = ["Merkmal", "weiblich", "divers", "männlich"] + AGES
    for i, h in enumerate(heads):
        hdr(ws.cell(top, 1 + i, h))
    vals = [fw, fd, fm] + fages
    cell(ws, f"A{top + 1}", "Anzahl", bold=True, bg=None, center=False)
    cell(ws, f"A{top + 2}", "Anteil", bold=True, bg=None, center=False)
    for i, f in enumerate(vals):
        c = 2 + i
        cell(ws, f"{L(c)}{top + 1}", f"={f}")
        den = f"SUM($B${top + 1}:$D${top + 1})" if i < 3 else f"SUM($E${top + 1}:$I${top + 1})"
        cell(ws, f"{L(c)}{top + 2}", f"=IF({den}=0,0,{L(c)}{top + 1}/{den})", bold=True, fmt="0%")


def pie_gender(ws, row, anchor, ttl):
    p = PieChart()
    p.title = ttl
    p.add_data(Reference(ws, min_col=2, max_col=4, min_row=row + 1), from_rows=True, titles_from_data=False)
    p.set_categories(Reference(ws, min_col=2, max_col=4, min_row=row))
    p.dataLabels = labels(True)
    color_points(p, GCOL)
    p.height, p.width = 7.5, 10
    ws.add_chart(p, anchor)


def bar_age(ws, row, anchor, ttl):
    b = BarChart()
    b.type = "col"
    b.title = ttl
    b.legend = None
    b.add_data(Reference(ws, min_col=5, max_col=9, min_row=row + 1), from_rows=True, titles_from_data=False)
    b.set_categories(Reference(ws, min_col=5, max_col=9, min_row=row))
    b.dataLabels = labels()
    color_points(b, ACOL)
    b.height, b.width = 7.5, 10
    ws.add_chart(b, anchor)


def bar_angebote(ws, top, first, last, anchor, ttl):
    c = BarChart()
    c.type = "bar"
    c.grouping = "stacked"
    c.overlap = 100
    c.gapWidth = 40
    c.title = ttl
    c.add_data(Reference(ws, min_col=4, max_col=6, min_row=top, max_row=last), titles_from_data=True)
    c.set_categories(Reference(ws, min_col=1, min_row=first, max_row=last))
    c.x_axis.scaling.orientation = "maxMin"
    color_series(c, GCOL)
    c.height, c.width = 18, 20
    ws.add_chart(c, anchor)


def setup_eval_sheet(ws, t, sub):
    ws.sheet_view.showGridLines = False
    title(ws, t, sub)
    ws.column_dimensions["A"].width = 18
    for i in range(2, 17):
        ws.column_dimensions[L(i)].width = 8.5
    ws.column_dimensions["Q"].width = 3
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


# =====================================================================
# WOCHE
# =====================================================================
wo = wb.create_sheet("WOCHE")
setup_eval_sheet(wo, "Wochenübersicht", "Kalenderwoche in der gelben Zelle eintragen – alles andere rechnet automatisch.")
wo["A1"] = f'="Wochenübersicht "&{ORG}'
cell(wo, "A3", "Kalenderwoche:", bold=True, bg=None, center=False, size=11)
KW_NOW = datetime.date.today().isocalendar()[1] if datetime.date.today().year == JAHR else 1
cell(wo, "B3", KW_NOW, bold=True, bg=YEL, size=12).protection = UNLOCK
cell(wo, "C3", "von", bg=None, size=10)
cell(wo, "D3", f"=DATE({Y},1,4)-WEEKDAY(DATE({Y},1,4),2)+1+($B$3-1)*7", bold=True, bg=None, fmt="DD.MM.YYYY", size=10)
wo.merge_cells("D3:E3")
cell(wo, "F3", "bis", bg=None, size=10)
cell(wo, "G3", "=D3+6", bold=True, bg=None, fmt="DD.MM.YYYY", size=10)
wo.merge_cells("G3:H3")
dvk = DataValidation(type="whole", operator="between", formula1="1", formula2="53", showErrorMessage=True,
                     error="Kalenderwoche 1–53")
dvk.add("B3")
wo.add_data_validation(dvk)
KWC = f",{R('C')},$B$3"
tiles(wo, 5, "COUNTIF($C$10:$C$16,\">0\")", "$C$17", "$D$17", "$E$17", "$F$17", "$G$17")

section(wo, "A8", "Tage der Woche")
DH = ["Tag", "Datum", "Einträge"] + [m[0] for m in MET]
for i, h in enumerate(DH):
    hdr(wo.cell(9, 1 + i, h))
for k in range(7):
    r = 10 + k
    wo[f"B{r}"] = f"=$D$3+{k}"
    wo[f"A{r}"] = f'=CHOOSE({k + 1},{WTAG})&" "&DAY(B{r})&"."&MONTH(B{r})&"."'
    wo[f"C{r}"] = f"=SUMIFS({K('F')},{K('A')},B{r})"
    for j, (name, ec, kc) in enumerate(MET):
        wo.cell(r, 4 + j, f"=SUMIFS({K(kc)},{K('A')},B{r})")
    for c in range(1, 14):
        x = wo.cell(r, c)
        x.border = BOX
        x.font = font(size=9)
        x.alignment = Alignment(horizontal="center")
    wo[f"B{r}"].number_format = "DD.MM."
    wo[f"D{r}"].number_format = "0.#"
wo["A17"] = "SUMME"
for c in range(3, 14):
    wo.cell(17, c, f"=SUM({L(c)}10:{L(c)}16)")
for c in range(1, 14):
    cell(wo, f"{L(c)}17", bold=True, bg=BLUE, center=c > 1)

section(wo, "A19", "Angebote dieser Woche")
wa1, wa2 = angebot_table(wo, 20, KWC)
section(wo, "A53", "Team diese Woche")
team_table(wo, 54, 1, KWC)
section(wo, "E53", "Wichtiges der Woche (aus Spalte 'Wichtig / Vorkommnisse')")
wichtig_list(wo, 54, 5, 25, "W", "AY")

ch = BarChart()
ch.type = "col"
ch.grouping = "stacked"
ch.overlap = 100
ch.gapWidth = 50
ch.title = "Kontakte pro Tag (w / d / m)"
ch.add_data(Reference(wo, min_col=5, max_col=7, min_row=9, max_row=16), titles_from_data=True)
ch.set_categories(Reference(wo, min_col=1, min_row=10, max_row=16))
color_series(ch, GCOL)
ch.height, ch.width = 8, 16
wo.add_chart(ch, "R5")
bar_angebote(wo, 20, wa1, wa2, "R22", "Kontakte pro Angebot – diese Woche")
protect(wo)

# =====================================================================
# MONAT
# =====================================================================
mo = wb.create_sheet("MONAT")
setup_eval_sheet(mo, "Monatsübersicht", "Monat in der gelben Zelle auswählen – alles andere rechnet automatisch.")
mo["A1"] = f'="Monatsübersicht "&{ORG}'
cell(mo, "A3", "Monat:", bold=True, bg=None, center=False, size=11)
M_NOW = MONATSNAMEN[datetime.date.today().month - 1] if datetime.date.today().year == JAHR else MONATSNAMEN[0]
mo.merge_cells("B3:C3")
cell(mo, "B3", M_NOW, bold=True, bg=YEL, size=12).protection = UNLOCK
cell(mo, "D3", "von", bg=None, size=10)
cell(mo, "E3", f"=DATE({Y},$AZ$3,1)", bold=True, bg=None, fmt="DD.MM.YYYY", size=10)
mo.merge_cells("E3:F3")
cell(mo, "G3", "bis", bg=None, size=10)
cell(mo, "H3", f"=DATE({Y},$AZ$3+1,0)", bold=True, bg=None, fmt="DD.MM.YYYY", size=10)
mo.merge_cells("H3:I3")
mo["AZ3"] = f"=IFERROR(MATCH($B$3,{MON},0),1)"
mo.column_dimensions["AZ"].hidden = True
dvm = DataValidation(type="list", formula1=MON, showErrorMessage=True, error="Bitte Monat aus der Liste wählen.")
dvm.add("B3")
mo.add_data_validation(dvm)
MC = f",{R('V')},$AZ$3"
km = lambda kc: f"SUMIFS({K(kc)},{K('D')},$AZ$3)"
tiles(mo, 5, km("G"), km("F"), km("H"), "$B$10", "$C$10", "$D$10")
section(mo, "A8", "Geschlecht & Alter")
gender_age_block(mo, 9, km("I"), km("J"), km("K"), [km(c) for c in "MNOPQ"])
section(mo, "A13", "Angebote im Monat")
ma1, ma2 = angebot_table(mo, 14, MC)
section(mo, "A47", "Team im Monat")
team_table(mo, 48, 1, MC)
section(mo, "E47", "Wichtiges im Monat")
wichtig_list(mo, 48, 5, 30, "X", "AY")
section(mo, "A80", "Kontakte pro Tag")
for i, h in enumerate(["Datum", "Tag", "Einträge", "Kontakte"]):
    hdr(mo.cell(81, 1 + i, h))
for d in range(31):
    r = 82 + d
    mo[f"A{r}"] = f'=IF({d + 1}>DAY($H$3),"",DATE({Y},$AZ$3,{d + 1}))'
    mo[f"B{r}"] = f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),{WTAG}))'
    mo[f"C{r}"] = f'=IF(A{r}="","",SUMIFS({K("F")},{K("A")},A{r}))'
    mo[f"D{r}"] = f'=IF(A{r}="","",SUMIFS({K("L")},{K("A")},A{r}))'
    for c in range(1, 5):
        x = mo.cell(r, c)
        x.border = BOX
        x.font = font(size=9)
        x.alignment = Alignment(horizontal="center")
    mo[f"A{r}"].number_format = "DD.MM."
pie_gender(mo, 9, "R5", "Geschlecht – Monat")
bar_age(mo, 9, "X5", "Altersgruppen – Monat")
bar_angebote(mo, 14, ma1, ma2, "R21", "Kontakte pro Angebot – Monat")
dc = BarChart()
dc.type = "col"
dc.title = "Kontakte pro Tag"
dc.legend = None
dc.gapWidth = 40
dc.add_data(Reference(mo, min_col=4, min_row=81, max_row=112), titles_from_data=True)
dc.set_categories(Reference(mo, min_col=1, min_row=82, max_row=112))
color_series(dc, ["2E75B6"])
dc.height, dc.width = 8, 20
mo.add_chart(dc, "R58")
protect(mo)

# =====================================================================
# JAHRESBERICHT
# =====================================================================
ja = wb.create_sheet("JAHRESBERICHT")
setup_eval_sheet(ja, "Jahresbericht", "Rechnet automatisch aus der Erfassung. Nur 'Vorjahr' (gelb) wird händisch eingetragen.")
ja["A1"] = f'="Jahresbericht "&{Y}&" – "&{ORG}'
kj = lambda kc: f"SUM({K(kc)})"
tiles(ja, 4, kj("G"), kj("F"), kj("H"), kj("I"), kj("J"), kj("K"))

section(ja, "A7", "Kennzahlen mit Vorjahresvergleich")
for i, h in enumerate(["Kennzahl", "", "Jahr", "Vorjahr", "Veränderung"]):
    hdr(ja.cell(8, 1 + i, h))
ja.merge_cells("A8:B8")
ja["C8"] = f"={Y}"
ja["D8"] = f'="Vorjahr ("&({Y}-1)&")"'
KZ = [("Öffnungstage", kj("G"), "0"), ("Einträge (Angebote)", kj("F"), "0"),
      ("Kontakte gesamt", kj("L"), "0"),
      ("Ø Kontakte pro Öffnungstag", f"IF({kj('G')}=0,0,{kj('L')}/{kj('G')})", "0.0"),
      ("Stunden", kj("H"), "0.#"), ("Kontakte weiblich", kj("I"), "0"), ("Kontakte divers", kj("J"), "0"),
      ("Kontakte männlich", kj("K"), "0")]
for k, (lab, f, fmt) in enumerate(KZ):
    r = 9 + k
    ja.merge_cells(f"A{r}:B{r}")
    cell(ja, f"A{r}", lab, bg=None, center=False)
    cell(ja, f"B{r}", bg=None)
    cell(ja, f"C{r}", f"={f}", bold=True, fmt=fmt)
    cell(ja, f"D{r}", None, bg=YEL, fmt=fmt).protection = UNLOCK
    cell(ja, f"E{r}", f'=IF(N(D{r})=0,"",C{r}/D{r}-1)', fmt='+0%;-0%;0%')

section(ja, "G7", "Textbaustein für den Jahresbericht")
ja.merge_cells("G8:P16")
T_OPEN, T_ENT, T_KON = kj("G"), kj("F"), kj("L")
T_W, T_D, T_M = kj("I"), kj("J"), kj("K")
ja["G8"] = (
    f'=IF({T_KON}=0,"Noch keine Einträge vorhanden – der Text entsteht automatisch, sobald erfasst wird.",'
    f'"Im Jahr "&{Y}&" hatte "&{ORG}&" an "&{T_OPEN}&" Tagen geöffnet. In "&{T_ENT}&" Angeboten wurden "'
    f'&{T_KON}&" Kontakte gezählt – das sind durchschnittlich "&ROUND({T_KON}/MAX(1,{T_OPEN}),1)&'
    f'" Kontakte pro Öffnungstag. "&ROUND({T_W}/{T_KON}*100,0)&" % der Kontakte waren weiblich, "'
    f'&ROUND({T_D}/{T_KON}*100,0)&" % divers und "&ROUND({T_M}/{T_KON}*100,0)&" % männlich. "'
    f'&"Die stärkste Altersgruppe war „"&INDEX($E$40:$I$40,MATCH(MAX($E$41:$I$41),$E$41:$I$41,0))&"“ ("'
    f'&ROUND(MAX($E$42:$I$42)*100,0)&" %). Das meistbesuchte Angebot war „"'
    f'&INDEX($A$46:$A$75,MATCH(MAX($G$46:$G$75),$G$46:$G$75,0))&"“ mit "&MAX($G$46:$G$75)&" Kontakten. "'
    f'&"Insgesamt wurden "&ROUND({kj("H")},1)&" Angebotsstunden dokumentiert.")')
ja["G8"].font = font(size=10)
ja["G8"].alignment = Alignment(wrap_text=True, vertical="top")
for r in range(8, 17):
    for c in range(7, 17):
        ja.cell(r, c).border = BOX
        ja.cell(r, c).fill = fill("FBFBFB")

section(ja, "A18", "Monate")
MH = ["Monat", "Öffnungstage", "Einträge"] + [m[0] for m in MET] + ["Ø Kontakte / Tag"]
for i, h in enumerate(MH):
    hdr(ja.cell(19, 1 + i, h))
for k in range(12):
    r = 20 + k
    ja[f"A{r}"] = f"=LISTEN!$H${2 + k}"
    ja[f"B{r}"] = f"=SUMIFS({K('G')},{K('D')},{k + 1})"
    ja[f"C{r}"] = f"=SUMIFS({K('F')},{K('D')},{k + 1})"
    for j, (name, ec, kc) in enumerate(MET):
        ja.cell(r, 4 + j, f"=SUMIFS({K(kc)},{K('D')},{k + 1})")
    ja[f"N{r}"] = f"=IF(B{r}=0,0,H{r}/B{r})"
    for c in range(1, 15):
        x = ja.cell(r, c)
        x.border = BOX
        x.font = font(size=9)
        if c > 1:
            x.alignment = Alignment(horizontal="center")
        if k % 2:
            x.fill = fill(ZEBRA)
    ja[f"D{r}"].number_format = "0.#"
    ja[f"N{r}"].number_format = "0.0"
ja["A32"] = "JAHR"
for c in range(2, 14):
    ja.cell(32, c, f"=SUM({L(c)}20:{L(c)}31)")
ja["N32"] = "=IF(B32=0,0,H32/B32)"
for c in range(1, 15):
    cell(ja, f"{L(c)}32", bold=True, bg=BLUE, center=c > 1)
ja["D32"].number_format = "0.#"
ja["N32"].number_format = "0.0"
ja["A33"] = "Kontrolle: Summe Monate = Summe Angebote?"
ja["A33"].font = font(size=8, italic=True)
ja["H33"] = '=IF(H32=G76,"✓ OK","⚠ prüfen: Einträge ohne Angebot?")'
ja["H33"].font = font(size=8, bold=True)

section(ja, "A38", "Geschlecht & Alter")
gender_age_block(ja, 40, "$E$32", "$F$32", "$G$32", ["$I$32", "$J$32", "$K$32", "$L$32", "$M$32"])
ja.row_dimensions[39].height = 4

section(ja, "A44", "Angebote im Jahr")
ja1, ja2 = angebot_table(ja, 45, "")

section(ja, "A78", "Wochentage")
for i, h in enumerate(["Wochentag", "Öffnungstage", "Kontakte", "Ø Kontakte"]):
    hdr(ja.cell(79, 1 + i, h))
for k, t in enumerate(["Montag", "Dienstag", "Mittwoch", "Donnerstag", "Freitag", "Samstag", "Sonntag"]):
    r = 80 + k
    cell(ja, f"A{r}", t, bg=None, center=False)
    cell(ja, f"B{r}", f"=SUMIFS({K('G')},{K('E')},{k + 1})", bg=None)
    cell(ja, f"C{r}", f"=SUMIFS({K('L')},{K('E')},{k + 1})", bg=None)
    cell(ja, f"D{r}", f"=IF(B{r}=0,0,C{r}/B{r})", bold=True, fmt="0.0")
section(ja, "F78", "Team im Jahr")
team_table(ja, 79, 6, "")

section(ja, "A111", "Kontakte pro Kalenderwoche")
for i, h in enumerate(["KW", "Öffnungstage", "Einträge", "w", "d", "m", "gesamt"]):
    hdr(ja.cell(112, 1 + i, h))
for k in range(53):
    r = 113 + k
    ja[f"A{r}"] = k + 1
    ja[f"B{r}"] = f"=SUMIFS({K('G')},{K('C')},A{r})"
    ja[f"C{r}"] = f"=SUMIFS({K('F')},{K('C')},A{r})"
    for j, kc in enumerate("IJKL"):
        ja.cell(r, 4 + j, f"=SUMIFS({K(kc)},{K('C')},A{r})")
    for c in range(1, 8):
        x = ja.cell(r, c)
        x.border = BOX
        x.font = font(size=9)
        x.alignment = Alignment(horizontal="center")
        if k % 2:
            x.fill = fill(ZEBRA)

mc = BarChart()
mc.type = "col"
mc.grouping = "stacked"
mc.overlap = 100
mc.gapWidth = 50
mc.title = "Kontakte pro Monat (w / d / m)"
mc.add_data(Reference(ja, min_col=5, max_col=7, min_row=19, max_row=31), titles_from_data=True)
mc.set_categories(Reference(ja, min_col=1, min_row=20, max_row=31))
color_series(mc, GCOL)
mc.height, mc.width = 8, 20
ja.add_chart(mc, "R4")
pie_gender(ja, 40, "R21", "Geschlecht – Jahr")
bar_age(ja, 40, "X21", "Altersgruppen – Jahr")
bar_angebote(ja, 45, ja1, ja2, "R37", "Kontakte pro Angebot – Jahr")
wd = BarChart()
wd.type = "col"
wd.title = "Ø Kontakte pro Öffnungstag nach Wochentag"
wd.legend = None
wd.add_data(Reference(ja, min_col=4, min_row=79, max_row=86), titles_from_data=True)
wd.set_categories(Reference(ja, min_col=1, min_row=80, max_row=86))
wd.dataLabels = labels()
color_series(wd, ["2E75B6"])
wd.height, wd.width = 8, 20
ja.add_chart(wd, "R74")
lk = LineChart()
lk.title = "Kontakte pro Kalenderwoche"
lk.add_data(Reference(ja, min_col=7, min_row=112, max_row=165), titles_from_data=True)
lk.set_categories(Reference(ja, min_col=1, min_row=113, max_row=165))
lk.legend = None
lk.series[0].smooth = False
color_series(lk, ["2E75B6"])
lk.height, lk.width = 8, 20
ja.add_chart(lk, "R91")
protect(ja)

# ---------- Reihenfolge, Start ----------
order = ["ANLEITUNG", "ERFASSUNG", "WOCHE", "MONAT", "JAHRESBERICHT", "LISTEN", "KALENDER"]
wb._sheets = [wb[n] for n in order]
wb["KALENDER"].sheet_properties.tabColor = "BFBFBF"
wb["LISTEN"].sheet_properties.tabColor = "FFD966"
wb["ERFASSUNG"].sheet_properties.tabColor = "70AD47"
for n in ("WOCHE", "MONAT", "JAHRESBERICHT"):
    wb[n].sheet_properties.tabColor = "2E75B6"
wb.active = 1
wb.calculation = CalcProperties(fullCalcOnLoad=True)

# ---------- optional: Einträge übernehmen ----------
if len(sys.argv) > 2:
    OUT = sys.argv[2]
    data = json.load(open(sys.argv[1], encoding="utf-8"))
    for i, n in enumerate(data.get("team", [])):
        li[f"C{2 + i}"] = n
    order_a = {x: i for i, x in enumerate(ANGEBOTE)}
    rows = sorted(data["eintraege"], key=lambda e: (e["datum"], order_a.get(e.get("angebot"), 99)))
    for r, e in enumerate(rows, FIRST):
        er[f"A{r}"] = datetime.date.fromisoformat(e["datum"])
        for key, col in [("angebot", "D"), ("stunden", "E"), ("w", "F"), ("d", "G"), ("m", "H"), ("u10", "J"),
                         ("ab10", "K"), ("ab14", "L"), ("ab16", "M"), ("ab18", "N"), ("dienst1", "P"),
                         ("dienst2", "Q"), ("dienst3", "R"), ("beschreibung", "S"), ("wichtig", "T")]:
            if e.get(key) not in (None, ""):
                er[f"{col}{r}"] = e[key]
        if not e.get("angebot"):
            er[f"D{r}"].fill = fill(YEL)

wb.save(OUT)
print("gespeichert:", OUT)
