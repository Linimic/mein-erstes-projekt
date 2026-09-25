# Erzeugt die Excel-Vorlage Dienstdoku_Vorlage_2026.xlsx (python3 erstelle_dienstdoku_vorlage.py)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule
from openpyxl.chart import BarChart, PieChart, LineChart, Reference
from openpyxl.chart.label import DataLabelList


from openpyxl.chart.series import DataPoint
GCOL = ["8E6CC4", "F2A541", "3E8EC2"]   # w, d, m
ACOL = ["9CC3E6", "5B9BD5", "2E75B6", "1F4E79"]  # ab10..ab18
def labels(pct=False):
    d = DataLabelList(); d.showVal = not pct; d.showPercent = pct
    d.showSerName = False; d.showCatName = False; d.showLegendKey = False; d.showLeaderLines = False
    return d
def color_series(ch, cols):
    for s, c in zip(ch.series, cols):
        s.graphicalProperties.solidFill = c; s.graphicalProperties.line.solidFill = c
def color_points(ch, cols):
    s = ch.series[0]
    for i, c in enumerate(cols):
        pt = DataPoint(idx=i); pt.graphicalProperties.solidFill = c; s.dPt.append(pt)

OUT = "Dienstdoku_Vorlage_2026.xlsx"
F = "Arial"
NAVY = "1F3864"; BLUE = "DDEBF7"; INPUT = "FFFFFF"; CALC = "F2F2F2"; YEL = "FFF2CC"
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
def font(**k): k.setdefault("name", F); return Font(**k)
def hdr(c, fill=NAVY):
    c.font = font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor=fill)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True); c.border = BOX

ANGEBOTE = ["MITTAG","JUGENDCAFE","JUGENDCAFE NEXT","OSA","VERMIETUNGEN","TRÄNKEWEG","TANZEN",
 "KONZERTE","JUNGS","MÄDCHEN","JSA","RADIKALISIERUNG","SCHULE","WORKSHOPS","CARITAS",
 "KINDERTHEATER","WOODROCK","KOOPERATIONEN","ÖA","OUTDOORFUN","ONLINE","MOJA","SKATERPLATZ","BERATUNG / TEL"]
NA = 30  # Plätze für Angebote
MONATE = ["JANUAR","FEBRUAR","MÄRZ","APRIL","MAI","JUNI","JULI","AUGUST","SEPTEMBER","OKTOBER","NOVEMBER","DEZEMBER"]
FIRST, LAST = 6, 405   # Datenzeilen Monatsblatt (400 Einträge)
SR = 6                 # erste Zeile Auswertung
SE = SR + NA - 1       # letzte Angebotszeile
ST = SE + 1            # Summenzeile

wb = Workbook()

# ---------------- Anleitung ----------------
a = wb.active; a.title = "ANLEITUNG"
a.sheet_view.showGridLines = False
a.column_dimensions["A"].width = 3; a.column_dimensions["B"].width = 26
for c in "CDEFGHIJKLMNOP": a.column_dimensions[c].width = 12
a["B2"] = "Dienstdokumentation – Anleitung"; a["B2"].font = font(size=16, bold=True, color=NAVY)
txt = [
 ("So funktioniert's", None),
 ("1.", "Im Blatt LISTEN das Jahr eintragen und bei Bedarf Angebote / Diensthabende anpassen (gelbe Zellen)."),
 ("2.", "Pro Monat gibt es EIN Blatt. Pro Angebot und Tag wird EINE Zeile ausgefüllt – nur die Angebote, die wirklich stattgefunden haben."),
 ("3.", "Weiße Zellen = Eingabe. Graue Zellen = rechnen automatisch (nicht überschreiben)."),
 ("4.", "Angebot und Diensthabende über das Dropdown wählen (Diensthabende kann auch frei getippt werden)."),
 ("5.", "Rechts neben der Tabelle steht sofort die Monatsauswertung mit Diagrammen (Geschlecht, Alter, Angebote)."),
 ("6.", "Das Blatt JAHRESAUSWERTUNG fasst alle Monate automatisch zusammen – inkl. Verlauf und Diagrammen."),
 ("7.", "Orange markiert: 'gesamt Alter' weicht von 'gesamt' ab → Zahlen prüfen. Wochenenden sind grau hinterlegt."),
 ("8.", "Mit dem Filter-Pfeil in der Kopfzeile lässt sich z.B. nach Angebot filtern – die Summenzeile oben rechnet dann nur die gefilterten Zeilen."),
]
r = 4
for k, v in txt:
    a.cell(r, 2, k).font = font(bold=(v is None), size=12 if v is None else 10, color=NAVY if v is None else "000000")
    if v: a.cell(r, 3, v).font = font(size=10)
    r += 1
r += 1
a.cell(r, 2, "Beispielzeile (so sieht ein Eintrag aus)").font = font(bold=True, size=12, color=NAVY); r += 1
heads = ["Datum","Tag","Angebot","Stunden","w","d","m","gesamt","ab 10","ab 14","ab 16","ab 18","gesamt Alter","Diensthabende","Beschreibung","Wichtig"]
for i, h in enumerate(heads): hdr(a.cell(r, 2 + i, h))
ex = ["06.01.2026","Di","JUGENDCAFE",4,7,1,12,20,5,9,4,2,20,"Anna, Max","Billard-Turnier","Fenster WC defekt"]
for i, v in enumerate(ex):
    c = a.cell(r + 1, 2 + i, v); c.font = font(size=10); c.border = BOX
a.column_dimensions["B"].width = 26
r += 3
a.cell(r, 2, "Spaltenbedeutung").font = font(bold=True, size=12, color=NAVY); r += 1
for k, v in [("w / d / m","Anzahl Besucher:innen weiblich / divers / männlich"),
             ("ab 10 … ab 18","Anzahl nach Altersgruppe (10–13, 14–15, 16–17, 18+)"),
             ("gesamt Alter","Summe der Altersgruppen (Kontrolle zu 'gesamt')"),
             ("Einträge","Wie oft das Angebot im Zeitraum stattgefunden hat")]:
    a.cell(r, 2, k).font = font(bold=True, size=10); a.cell(r, 3, v).font = font(size=10); r += 1

# ---------------- Listen ----------------
li = wb.create_sheet("LISTEN")
li.sheet_view.showGridLines = False
li["A1"] = "Angebote"; li["C1"] = "Diensthabende"; li["E1"] = "Jahr"
for c in ("A1","C1","E1"): hdr(li[c])
li.column_dimensions["A"].width = 26; li.column_dimensions["C"].width = 22; li.column_dimensions["E"].width = 10
li["E2"] = 2026
for i in range(NA):
    c = li.cell(2 + i, 1, ANGEBOTE[i] if i < len(ANGEBOTE) else None)
    c.fill = PatternFill("solid", fgColor=YEL); c.border = BOX; c.font = font()
for i in range(30):
    c = li.cell(2 + i, 3); c.fill = PatternFill("solid", fgColor=YEL); c.border = BOX; c.font = font()
li["E2"].fill = PatternFill("solid", fgColor=YEL); li["E2"].border = BOX; li["E2"].font = font(bold=True)
li["G2"] = "Gelbe Zellen anpassen. Neue Angebote einfach in freie Zeilen schreiben – sie erscheinen automatisch in Dropdown & Auswertung."
li["G2"].font = font(italic=True, size=9, color="595959")
ANG_RANGE = f"LISTEN!$A$2:$A${NA+1}"

# ---------------- Monatsblätter ----------------
COLS = heads  # A..P
WIDTH = [11,5,20,8,6,6,6,8,7,7,7,7,8,16,32,24]
SUMCOLS = ["Stunden","w","d","m","gesamt","ab 10","ab 14","ab 16","ab 18"]
SRC = {"Stunden":"D","w":"E","d":"F","m":"G","gesamt":"H","ab 10":"I","ab 14":"J","ab 16":"K","ab 18":"L"}

def month_sheet(mi, name):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(WIDTH): ws.column_dimensions[L(i + 1)].width = w
    ws.column_dimensions["Q"].width = 3
    ws["A1"] = f'="{name} "&LISTEN!$E$2'; ws["A1"].font = font(size=16, bold=True, color=NAVY)
    ws["A2"] = "Eine Zeile pro Angebot und Tag. Weiße Zellen ausfüllen, graue rechnen automatisch."
    ws["A2"].font = font(italic=True, size=9, color="595959")
    # Summenzeile (Zeile 4, filterfähig)
    ws["C4"] = "SUMME (gefiltert)"; ws["C4"].font = font(bold=True); ws["C4"].alignment = Alignment(horizontal="right")
    for col in "DEFGHIJKLM":
        c = ws[f"{col}4"]; c.value = f"=SUBTOTAL(9,{col}{FIRST}:{col}{LAST})"
        c.font = font(bold=True); c.fill = PatternFill("solid", fgColor=BLUE); c.border = BOX
        c.alignment = Alignment(horizontal="center")
    for i, h in enumerate(COLS): hdr(ws.cell(5, i + 1, h))
    ws.row_dimensions[5].height = 30
    y = "LISTEN!$E$2"
    for r in range(FIRST, LAST + 1):
        ws[f"B{r}"] = f'=IF(A{r}="","",CHOOSE(WEEKDAY(A{r},2),"Mo","Di","Mi","Do","Fr","Sa","So"))'
        ws[f"H{r}"] = f'=IF(COUNT(E{r}:G{r})=0,"",SUM(E{r}:G{r}))'
        ws[f"M{r}"] = f'=IF(COUNT(I{r}:L{r})=0,"",SUM(I{r}:L{r}))'
        for col in range(1, 17):
            c = ws.cell(r, col); c.border = BOX; c.font = font(size=10)
            if L(col) in "BHM": c.fill = PatternFill("solid", fgColor=CALC)
            if L(col) in "ABDEFGHIJKLM": c.alignment = Alignment(horizontal="center")
            if L(col) in "OP": c.alignment = Alignment(wrap_text=True, vertical="top")
        ws[f"A{r}"].number_format = "DD.MM.YYYY"
        ws[f"D{r}"].number_format = "0.0#"
    ws.freeze_panes = f"A{FIRST}"
    ws.auto_filter.ref = f"A5:P{LAST}"
    # Validierungen
    dv = DataValidation(type="date", operator="between",
        formula1=f"DATE({y},{mi+1},1)", formula2=f"DATE({y},{mi+2},0)" if mi < 11 else f"DATE({y},12,31)",
        showErrorMessage=True, errorTitle="Datum", error=f"Bitte ein Datum im {name} eingeben (TT.MM.JJJJ).")
    dv.add(f"A{FIRST}:A{LAST}"); ws.add_data_validation(dv)
    dv2 = DataValidation(type="list", formula1=ANG_RANGE, showErrorMessage=True, errorTitle="Angebot",
        error="Bitte ein Angebot aus der Liste wählen (neue Angebote im Blatt LISTEN ergänzen).")
    dv2.add(f"C{FIRST}:C{LAST}"); ws.add_data_validation(dv2)
    dv3 = DataValidation(type="list", formula1="LISTEN!$C$2:$C$31", showErrorMessage=False)
    dv3.add(f"N{FIRST}:N{LAST}"); ws.add_data_validation(dv3)
    dv4 = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", showErrorMessage=True,
        error="Bitte eine ganze Zahl ≥ 0 eingeben.")
    dv4.add(f"E{FIRST}:G{LAST}"); dv4.add(f"I{FIRST}:L{LAST}"); ws.add_data_validation(dv4)
    dv5 = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", showErrorMessage=True,
        error="Stunden zwischen 0 und 24.")
    dv5.add(f"D{FIRST}:D{LAST}"); ws.add_data_validation(dv5)
    # bedingte Formatierung
    ws.conditional_formatting.add(f"H{FIRST}:H{LAST}",
        FormulaRule(formula=[f'AND(ISNUMBER(M{FIRST}),ISNUMBER(H{FIRST}),M{FIRST}<>H{FIRST})'],
                    fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"M{FIRST}:M{LAST}",
        FormulaRule(formula=[f'AND(ISNUMBER(M{FIRST}),ISNUMBER(H{FIRST}),M{FIRST}<>H{FIRST})'],
                    fill=PatternFill("solid", fgColor="F8CBAD")))
    ws.conditional_formatting.add(f"A{FIRST}:A{LAST}",
        FormulaRule(formula=[f'AND(ISNUMBER(A{FIRST}),WEEKDAY(A{FIRST},2)>5)'], fill=PatternFill("solid", fgColor="D9D9D9")))

    # ---- Monatsauswertung ab Spalte R ----
    base = 18  # R
    ws.cell(1, base, "Monatsauswertung").font = font(size=14, bold=True, color=NAVY)
    ws.cell(2, base, "rechnet automatisch – nicht überschreiben").font = font(italic=True, size=9, color="595959")
    sh = ["Angebot"] + SUMCOLS + ["Einträge"]
    for i, h in enumerate(sh): hdr(ws.cell(SR - 1, base + i, h))
    ws.column_dimensions[L(base)].width = 20
    for i in range(1, len(sh)): ws.column_dimensions[L(base + i)].width = 8
    for k in range(NA):
        r = SR + k
        ws.cell(r, base, f'=IF(LISTEN!$A${2+k}="","",LISTEN!$A${2+k})')
        for i, h in enumerate(SUMCOLS):
            s = SRC[h]
            ws.cell(r, base + 1 + i, f'=IF(${L(base)}{r}="","",SUMIFS(${s}${FIRST}:${s}${LAST},$C${FIRST}:$C${LAST},${L(base)}{r}))')
        ws.cell(r, base + 10, f'=IF(${L(base)}{r}="","",COUNTIF($C${FIRST}:$C${LAST},${L(base)}{r}))')
        for i in range(len(sh)):
            c = ws.cell(r, base + i); c.border = BOX; c.font = font(size=9); c.fill = PatternFill("solid", fgColor=CALC)
            if i: c.alignment = Alignment(horizontal="center")
    ws.cell(ST, base, "SUMME")
    for i in range(1, len(sh)):
        col = L(base + i); ws.cell(ST, base + i, f"=SUM({col}{SR}:{col}{SE})")
    for i in range(len(sh)):
        c = ws.cell(ST, base + i); c.font = font(bold=True, size=9); c.fill = PatternFill("solid", fgColor=BLUE); c.border = BOX
        if i: c.alignment = Alignment(horizontal="center")
    # Kennzahlen
    kr = ST + 2
    ws.cell(kr, base, "Kennzahlen").font = font(bold=True, color=NAVY)
    kz = [("Öffnungstage (Tage mit Einträgen)", f'=SUMPRODUCT((A{FIRST}:A{LAST}<>"")/COUNTIF(A{FIRST}:A{LAST},A{FIRST}:A{LAST}&""))'),
          ("Ø Besucher:innen pro Öffnungstag", f'=IF({L(base)}{kr+1}=0,0,{L(base+5)}{ST}/{L(base)}{kr+1})'),
          ("Anteil weiblich", f'=IF({L(base+5)}{ST}=0,0,{L(base+2)}{ST}/{L(base+5)}{ST})'),
          ("Anteil divers", f'=IF({L(base+5)}{ST}=0,0,{L(base+3)}{ST}/{L(base+5)}{ST})'),
          ("Anteil männlich", f'=IF({L(base+5)}{ST}=0,0,{L(base+4)}{ST}/{L(base+5)}{ST})')]
    # Wert in Spalte base+4 (neben Beschriftung)
    for j, (lab, f) in enumerate(kz):
        rr = kr + 1 + j
        ws.cell(rr, base, lab).font = font(size=9)
        f = f.replace(f"{L(base)}{kr+1}", f"{L(base+4)}{kr+1}")
        c = ws.cell(rr, base + 4, f); c.font = font(bold=True, size=9); c.border = BOX; c.fill = PatternFill("solid", fgColor=CALC)
        c.alignment = Alignment(horizontal="center")
        c.number_format = "0.0%" if "Anteil" in lab else ("0.0" if "Ø" in lab else "0")
    # Diagramme
    add_charts(ws, base, ST, SR, SE, anchor_row=kr + 8, title_suffix=name.capitalize())
    return ws

def add_charts(ws, base, st, sr, se, anchor_row, title_suffix):
    # Geschlecht (Torte)
    p = PieChart(); p.title = f"Geschlecht – {title_suffix}"
    p.add_data(Reference(ws, min_col=base + 2, max_col=base + 4, min_row=st), from_rows=True, titles_from_data=False)
    p.set_categories(Reference(ws, min_col=base + 2, max_col=base + 4, min_row=sr - 1))
    p.dataLabels = labels(True); color_points(p, GCOL)
    p.height, p.width = 7, 11
    ws.add_chart(p, f"{L(base)}{anchor_row}")
    # Alter (Säulen)
    b = BarChart(); b.type = "col"; b.title = f"Altersgruppen – {title_suffix}"; b.legend = None
    b.add_data(Reference(ws, min_col=base + 6, max_col=base + 9, min_row=st), from_rows=True, titles_from_data=False)
    b.set_categories(Reference(ws, min_col=base + 6, max_col=base + 9, min_row=sr - 1))
    b.dataLabels = labels(); color_points(b, ACOL)
    b.height, b.width = 7, 11
    ws.add_chart(b, f"{L(base + 7)}{anchor_row}")
    # Angebote (Balken, gestapelt nach Geschlecht)
    c = BarChart(); c.type = "bar"; c.grouping = "stacked"; c.overlap = 100
    c.title = f"Besucher:innen pro Angebot – {title_suffix}"
    c.add_data(Reference(ws, min_col=base + 2, max_col=base + 4, min_row=sr - 1, max_row=se), titles_from_data=True)
    c.set_categories(Reference(ws, min_col=base, min_row=sr, max_row=se))
    c.y_axis.scaling.orientation = "minMax"; c.x_axis.scaling.orientation = "maxMin"
    c.gapWidth = 40; color_series(c, GCOL); c.height, c.width = 18, 20
    ws.add_chart(c, f"{L(base)}{anchor_row + 15}")

for mi, m in enumerate(MONATE): month_sheet(mi, m)

# ---------------- Jahresauswertung ----------------
j = wb.create_sheet("JAHRESAUSWERTUNG", 2)
j.sheet_view.showGridLines = False
j["A1"] = '="JAHRESAUSWERTUNG "&LISTEN!$E$2'; j["A1"].font = font(size=16, bold=True, color=NAVY)
j["A2"] = "rechnet automatisch aus allen Monatsblättern – nichts eintragen"; j["A2"].font = font(italic=True, size=9, color="595959")
j.column_dimensions["A"].width = 20
for i in range(2, 16): j.column_dimensions[L(i)].width = 9
mh = ["Monat"] + SUMCOLS + ["Einträge", "Öffnungstage"]
j["A4"] = "Verlauf nach Monaten"; j["A4"].font = font(bold=True, size=12, color=NAVY)
for i, h in enumerate(mh): hdr(j.cell(5, 1 + i, h))
base = 18
kr = ST + 2
for k, m in enumerate(MONATE):
    r = 6 + k
    j.cell(r, 1, m.capitalize())
    for i in range(len(SUMCOLS) + 1):
        j.cell(r, 2 + i, f"='{m}'!{L(base + 1 + i)}{ST}")
    j.cell(r, 12, f"='{m}'!{L(base + 4)}{kr + 1}")
    for i in range(len(mh)):
        c = j.cell(r, 1 + i); c.border = BOX; c.font = font(size=10)
        if i: c.alignment = Alignment(horizontal="center")
        if i and k % 2: c.fill = PatternFill("solid", fgColor="F7F9FC")
j.cell(18, 1, "JAHR")
for i in range(1, len(mh)): j.cell(18, 1 + i, f"=SUM({L(1+i)}6:{L(1+i)}17)")
for i in range(len(mh)):
    c = j.cell(18, 1 + i); c.font = font(bold=True); c.fill = PatternFill("solid", fgColor=BLUE); c.border = BOX
    if i: c.alignment = Alignment(horizontal="center")
# Anteile
j["A20"] = "Anteile Jahr"; j["A20"].font = font(bold=True, size=12, color=NAVY)
anteile = [("weiblich","C"),("divers","D"),("männlich","E"),("ab 10","G"),("ab 14","H"),("ab 16","I"),("ab 18","J")]
for i, (lab, col) in enumerate(anteile):
    hdr(j.cell(21, 1 + i if i else 1, lab)) if False else None
j.cell(21, 1, "Merkmal"); j.cell(22, 1, "Anzahl"); j.cell(23, 1, "Anteil")
hdr(j["A21"])
for i, (lab, col) in enumerate(anteile):
    hdr(j.cell(21, 2 + i, lab))
    c = j.cell(22, 2 + i, f"={col}18"); c.border = BOX; c.font = font(); c.alignment = Alignment(horizontal="center")
    den = "$F$18" if i < 3 else "SUM($G$18:$J$18)"
    c = j.cell(23, 2 + i, f"=IF({den}=0,0,{col}18/{den})"); c.number_format = "0.0%"
    c.border = BOX; c.font = font(bold=True); c.alignment = Alignment(horizontal="center")
for rr in (22, 23): j.cell(rr, 1).font = font(bold=True); j.cell(rr, 1).border = BOX
j["J22"] = "Ø Besucher:innen / Öffnungstag"; j["J22"].font = font(size=9)
j["M22"] = "=IF(L18=0,0,F18/L18)"; j["M22"].number_format = "0.0"; j["M22"].font = font(bold=True); j["M22"].border = BOX

# Angebote Jahr
AR = 26
j.cell(AR - 1, 1, "Angebote im Jahr").font = font(bold=True, size=12, color=NAVY)
ah = ["Angebot"] + SUMCOLS + ["Einträge", "Anteil Besuche"]
for i, h in enumerate(ah): hdr(j.cell(AR, 1 + i, h))
for k in range(NA):
    r = AR + 1 + k
    j.cell(r, 1, f'=IF(LISTEN!$A${2+k}="","",LISTEN!$A${2+k})')
    for i in range(len(SUMCOLS) + 1):
        col = L(base + 1 + i); rr = SR + k
        refs = "+".join(f"N('{m}'!{col}{rr})" for m in MONATE)
        j.cell(r, 2 + i, f'=IF($A{r}="","",{refs})')
    j.cell(r, 12, f'=IF(OR($A{r}="",$F${AR+NA+1}=0),"",F{r}/$F${AR+NA+1})').number_format = "0.0%"
    for i in range(len(ah)):
        c = j.cell(r, 1 + i); c.border = BOX; c.font = font(size=10)
        if i: c.alignment = Alignment(horizontal="center")
        if k % 2: c.fill = PatternFill("solid", fgColor="F7F9FC")
TR = AR + NA + 1
j.cell(TR, 1, "SUMME")
for i in range(1, 11): j.cell(TR, 1 + i, f"=SUM({L(1+i)}{AR+1}:{L(1+i)}{AR+NA})")
for i in range(len(ah)):
    c = j.cell(TR, 1 + i); c.font = font(bold=True); c.fill = PatternFill("solid", fgColor=BLUE); c.border = BOX
    if i: c.alignment = Alignment(horizontal="center")
j.cell(TR + 1, 1, "Kontrolle: Summe Angebote = Summe Monate?").font = font(size=9, italic=True)
j.cell(TR + 1, 6, f'=IF(F{TR}=F18,"✓ OK","⚠ prüfen")').font = font(size=9, bold=True)
j.freeze_panes = "B6"

# Jahresdiagramme (rechts, ab Spalte O)
lc = BarChart(); lc.type = "col"; lc.grouping = "stacked"; lc.overlap = 100
lc.title = "Besucher:innen pro Monat nach Geschlecht"
lc.add_data(Reference(j, min_col=3, max_col=5, min_row=5, max_row=17), titles_from_data=True)
lc.set_categories(Reference(j, min_col=1, min_row=6, max_row=17))
lc.gapWidth = 50; color_series(lc, GCOL); lc.height, lc.width = 8, 20
j.add_chart(lc, "O4")
p = PieChart(); p.title = "Geschlecht – Jahr"
p.add_data(Reference(j, min_col=2, max_col=4, min_row=22), from_rows=True, titles_from_data=False)
p.set_categories(Reference(j, min_col=2, max_col=4, min_row=21))
p.dataLabels = labels(True); color_points(p, GCOL); p.height, p.width = 8, 11
j.add_chart(p, "O21")
b = BarChart(); b.type = "col"; b.title = "Altersgruppen – Jahr"; b.legend = None
b.add_data(Reference(j, min_col=5, max_col=8, min_row=22), from_rows=True, titles_from_data=False)
b.set_categories(Reference(j, min_col=5, max_col=8, min_row=21))
b.dataLabels = labels(); color_points(b, ACOL); b.height, b.width = 8, 11
j.add_chart(b, "V21")
ag = LineChart(); ag.title = "Altersgruppen im Verlauf"
ag.add_data(Reference(j, min_col=7, max_col=10, min_row=5, max_row=17), titles_from_data=True)
ag.set_categories(Reference(j, min_col=1, min_row=6, max_row=17)); ag.height, ag.width = 8, 20
for sr_ in ag.series: sr_.smooth = False
color_series(ag, ACOL)
j.add_chart(ag, "O38")
c = BarChart(); c.type = "bar"; c.grouping = "stacked"; c.overlap = 100
c.title = "Besucher:innen pro Angebot – Jahr"
c.add_data(Reference(j, min_col=3, max_col=5, min_row=AR, max_row=AR + NA), titles_from_data=True)
c.set_categories(Reference(j, min_col=1, min_row=AR + 1, max_row=AR + NA))
c.x_axis.scaling.orientation = "maxMin"; c.gapWidth = 40; color_series(c, GCOL); c.height, c.width = 18, 20
j.add_chart(c, "O55")
h = BarChart(); h.type = "bar"; h.title = "Stunden pro Angebot – Jahr"; h.legend = None
h.add_data(Reference(j, min_col=2, min_row=AR, max_row=AR + NA), titles_from_data=True)
h.set_categories(Reference(j, min_col=1, min_row=AR + 1, max_row=AR + NA))
h.x_axis.scaling.orientation = "maxMin"; h.gapWidth = 40; color_series(h, ["2E75B6"]); h.height, h.width = 18, 20
j.add_chart(h, "A62")

# Reihenfolge: Anleitung, Jahresauswertung, Monate, Listen
wb.move_sheet("LISTEN", offset=len(wb.sheetnames))
wb.active = 1

# Optional: Einträge übernehmen  (python3 erstelle_dienstdoku_vorlage.py daten.json ziel.xlsx)
import sys, json, datetime
if len(sys.argv) > 2:
    OUT = sys.argv[2]
    eintraege = json.load(open(sys.argv[1], encoding="utf-8"))
    nextrow = {m: FIRST for m in MONATE}
    order = {a: i for i, a in enumerate(ANGEBOTE)}
    eintraege.sort(key=lambda e: (e["datum"], order.get(e.get("angebot"), 99)))
    for e in eintraege:
        d = datetime.date.fromisoformat(e["datum"]); m = MONATE[d.month - 1]
        ws = wb[m]; r = nextrow[m]; nextrow[m] += 1
        ws[f"A{r}"] = d
        for key, col in [("angebot","C"),("stunden","D"),("w","E"),("d","F"),("m","G"),("ab10","I"),("ab14","J"),
                         ("ab16","K"),("ab18","L"),("dienst","N"),("beschreibung","O"),("wichtig","P")]:
            if e.get(key) not in (None, ""): ws[f"{col}{r}"] = e[key]
        if not e.get("angebot"):
            ws[f"C{r}"].fill = PatternFill("solid", fgColor=YEL)
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT); print("ok", wb.sheetnames)
